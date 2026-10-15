"""All four export formats must return valid, parseable files."""
from __future__ import annotations

import io
import zipfile

import pytest

pytestmark = pytest.mark.asyncio


@pytest.fixture
async def analyzed(client, guest, dataset):
    r = await client.post("/api/analyze", json={"dataset_id": dataset["dataset_id"]},
                          headers=guest["headers"])
    assert r.status_code == 200, r.text
    return guest, dataset, r.json()


async def test_pptx_export_is_a_real_presentation(client, analyzed):
    guest, dataset, _plan = analyzed
    r = await client.get(f"/api/export/{dataset['dataset_id']}?format=pptx",
                         headers=guest["headers"])
    assert r.status_code == 200
    assert "presentationml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    from pptx import Presentation
    prs = Presentation(io.BytesIO(r.content))
    assert len(prs.slides) >= 5


async def test_xlsx_export_opens_in_openpyxl(client, analyzed):
    guest, dataset, _plan = analyzed
    r = await client.get(f"/api/export/{dataset['dataset_id']}?format=xlsx",
                         headers=guest["headers"])
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames
    assert zipfile.is_zipfile(io.BytesIO(r.content))


async def test_pdf_export_is_a_pdf(client, analyzed):
    guest, dataset, _plan = analyzed
    for path in (f"/api/export/{dataset['dataset_id']}?format=pdf",
                 f"/api/export/pdf/{dataset['dataset_id']}"):
        r = await client.get(path, headers=guest["headers"])
        assert r.status_code == 200, path
        assert r.content.startswith(b"%PDF")
        assert r.content.rstrip().endswith(b"%%EOF")


async def test_png_export_per_slide(client, analyzed):
    guest, dataset, plan = analyzed
    for slide in plan["slides"][:4]:
        r = await client.get(f"/api/export/png/{dataset['dataset_id']}/{slide['id']}",
                             headers=guest["headers"])
        assert r.status_code == 200, slide["id"]
        assert r.content.startswith(b"\x89PNG"), slide["type"]


async def test_png_unknown_slide_is_404(client, analyzed):
    guest, dataset, _plan = analyzed
    r = await client.get(f"/api/export/png/{dataset['dataset_id']}/does-not-exist",
                         headers=guest["headers"])
    assert r.status_code == 404


async def test_bad_format_rejected(client, analyzed):
    guest, dataset, _plan = analyzed
    r = await client.get(f"/api/export/{dataset['dataset_id']}?format=ppt",
                         headers=guest["headers"])
    assert r.status_code == 400


def test_renderer_covers_every_chart_type():
    """render_slide must emit a PNG for all 37 chart types (PNG export path)."""
    import numpy as np
    import pandas as pd

    import charting
    import planner
    from exporters import render_slide
    from tests.test_charts import _spec_for  # noqa: PLC0415

    rng = np.random.default_rng(3)
    n = 120
    df = pd.DataFrame({
        "date": pd.date_range("2024-01-01", periods=n, freq="D"),
        "region": rng.choice(["North", "South", "East", "West"], n),
        "channel": rng.choice(["Retail", "Online"], n),
        "revenue": rng.gamma(4, 200, n).round(2),
        "units": rng.integers(1, 40, n),
        "margin": rng.normal(0.3, 0.04, n).round(3),
        "note": [f"words about row {i}" for i in range(n)],
    })
    types = {"date": "datetime", "region": "categorical", "channel": "categorical",
             "revenue": "numeric", "units": "numeric", "margin": "numeric", "note": "text"}
    for t in sorted(charting.BUILDERS):
        spec = _spec_for(t)
        data = charting.build(df, types, dict(spec))
        slide = {"id": t, "type": t, "title": planner.title_for(spec, "S"), "spec": spec,
                 "data": data, "insight": ""}
        png = render_slide(slide)
        assert png.startswith(b"\x89PNG"), t
